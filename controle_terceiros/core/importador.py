"""Lê a planilha Excel antiga (abas ENVIO, RECEBIMENTO e abas por fornecedor)
e converte para o formato do sistema."""
from __future__ import annotations

import datetime as dt
import re

import pandas as pd
from openpyxl import load_workbook

from core.schema import PRODUTO_PADRAO, SERV_CANECA, SERV_PERFIL

ABAS_FORNECEDOR_PERFIL = {"UNIFORJA", "MNG"}


def _txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _num(v):
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return 0


def _nf_op(texto: str):
    """'162943/op954881' -> ('162943', '954881')."""
    texto = _txt(texto)
    m = re.search(r"op\s*(\d+)", texto, flags=re.I)
    op = m.group(1) if m else ""
    nf = re.sub(r"/?\s*op\s*\d+", "", texto, flags=re.I).strip(" /")
    return nf, op


def ler_planilha_antiga(arquivo) -> dict[str, pd.DataFrame]:
    wb = load_workbook(arquivo, data_only=True, read_only=True)
    envios, recebs, prog = [], [], []

    if "ENVIO" in wb.sheetnames:
        for row in wb["ENVIO"].iter_rows(min_row=6, values_only=True):
            row = list(row) + [None] * 8
            data, tipo, forn, qtd, doc, resp = row[1], row[3], row[4], row[5], row[6], row[7]
            if not isinstance(data, dt.datetime) or not _num(qtd):
                continue
            nf, op = _nf_op(doc)
            envios.append({"data": data.date(), "fornecedor": _txt(forn).upper(),
                           "servico": SERV_CANECA if "CANECA" in _txt(tipo).upper() else SERV_PERFIL,
                           "produto": PRODUTO_PADRAO, "quantidade": _num(qtd), "nf": nf, "op": op,
                           "corrida": "", "responsavel": _txt(resp).title(),
                           "observacao": f"Importado da planilha ({_txt(doc)})"})

    serv_forn = {}
    for e in envios:
        serv_forn.setdefault(e["fornecedor"], e["servico"])

    if "RECEBIMENTO" in wb.sheetnames:
        for row in wb["RECEBIMENTO"].iter_rows(min_row=6, values_only=True):
            row = list(row) + [None] * 11
            data, forn, aprov, sucata, nf, resp = row[1], row[3], row[4], row[7], row[9], row[10]
            if not isinstance(data, dt.datetime) or not (_num(aprov) or _num(sucata)):
                continue
            f = _txt(forn).upper()
            recebs.append({"data": data.date(), "fornecedor": f,
                           "servico": serv_forn.get(f, SERV_PERFIL if f in ABAS_FORNECEDOR_PERFIL else SERV_CANECA),
                           "produto": PRODUTO_PADRAO, "qtd_aprovada": _num(aprov), "qtd_nao_conforme": 0,
                           "qtd_sucata": _num(sucata), "motivo_nc": "", "nf_retorno": _txt(nf),
                           "nf_referencia": "", "responsavel": _txt(resp).title(),
                           "observacao": "Importado da planilha"})

    for nome in wb.sheetnames:
        if nome in ("MENU", "INDICADORES", "ENVIO", "RECEBIMENTO"):
            continue
        forn = nome.strip().upper()
        servico = SERV_PERFIL if forn in ABAS_FORNECEDOR_PERFIL else SERV_CANECA
        for row in wb[nome].iter_rows(values_only=True):
            row = list(row) + [None] * 20
            for col_data, col_prev in ((1, 3), (10, 12)):   # colunas B/D e K/M
                d, p = row[col_data], row[col_prev]
                if isinstance(d, dt.datetime) and _num(p) > 0:
                    prog.append({"data": d.date(), "fornecedor": forn, "servico": servico,
                                 "produto": PRODUTO_PADRAO, "qtd_prevista": _num(p)})

    return {"envios": pd.DataFrame(envios), "recebimentos": pd.DataFrame(recebs),
            "programacao": pd.DataFrame(prog)}


def remover_duplicados(novos: pd.DataFrame, existentes: pd.DataFrame, campos: list[str]) -> pd.DataFrame:
    """Evita importar duas vezes o mesmo registro."""
    if novos.empty or existentes.empty:
        return novos

    def chave(df):
        d = df[campos].copy()
        d["data"] = pd.to_datetime(d["data"]).dt.strftime("%Y-%m-%d")  # ambos já são datas
        return d.astype(str).agg("|".join, axis=1)

    return novos[~chave(novos).isin(set(chave(existentes)))]
